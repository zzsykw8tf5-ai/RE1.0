import io
from typing import Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False


def _safe_value(val):
    """Convert cell value to a safe Python type."""
    if val is None:
        return None
    if hasattr(val, 'isoformat'):
        return val.isoformat()
    return val


def parse_property_master(file_bytes: bytes) -> dict:
    """
    Parse the 'Objektstammdaten' sheet from an uploaded Excel file.
    Returns a dict with property fields.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    sheet_name = None
    for name in wb.sheetnames:
        if "stammdaten" in name.lower() or "objekt" in name.lower():
            sheet_name = name
            break
    if sheet_name is None and wb.sheetnames:
        sheet_name = wb.sheetnames[0]

    if sheet_name is None:
        return {}

    ws = wb[sheet_name]
    data = {}

    field_map = {
        "name": ["name", "objektname", "bezeichnung"],
        "address": ["adresse", "address", "straße", "strasse", "anschrift"],
        "city": ["stadt", "city", "ort"],
        "zip_code": ["plz", "postleitzahl", "zip", "zip_code"],
        "property_type": ["objekttyp", "property_type", "nutzungsart", "typ"],
        "construction_year": ["baujahr", "construction_year", "erbaut"],
        "total_area_sqm": ["gesamtfläche", "total_area", "mietfläche", "fläche", "gesamtflaeche"],
        "land_area_sqm": ["grundstücksfläche", "land_area", "grundstücksflaeche"],
        "floors": ["etagen", "geschosse", "floors", "stockwerke"],
        "units": ["einheiten", "units", "wohnungen", "mieteinheiten"],
        "purchase_price": ["kaufpreis", "purchase_price", "erwerbspreis"],
        "purchase_date": ["kaufdatum", "purchase_date", "erwerbsdatum"],
    }

    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 2:
            continue
        key_cell = str(row[0]).strip().lower() if row[0] else ""
        value_cell = row[1]

        for field, synonyms in field_map.items():
            if any(syn in key_cell for syn in synonyms):
                data[field] = _safe_value(value_cell)
                break

    return data


def parse_tenant_list(file_bytes: bytes) -> list:
    """
    Parse the 'Mieterliste' sheet.
    Returns list of tenant dicts.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    sheet_name = None
    for name in wb.sheetnames:
        if "mieter" in name.lower() or "tenant" in name.lower():
            sheet_name = name
            break
    if sheet_name is None and len(wb.sheetnames) > 1:
        sheet_name = wb.sheetnames[1]
    elif sheet_name is None and wb.sheetnames:
        sheet_name = wb.sheetnames[0]

    if sheet_name is None:
        return []

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    # First non-empty row is header
    header_row = None
    header_idx = 0
    for i, row in enumerate(rows):
        if any(cell is not None for cell in row):
            header_row = [str(c).strip().lower() if c else "" for c in row]
            header_idx = i
            break

    if header_row is None:
        return []

    col_map = {
        "name": ["name", "mietername", "mieter"],
        "unit": ["einheit", "unit", "wohnung", "nr"],
        "area_sqm": ["fläche", "flaeche", "area", "sqm", "m2", "qm"],
        "monthly_rent": ["monatsmiete", "monthly_rent", "kaltmiete_monat", "miete/monat"],
        "annual_rent": ["jahresmiete", "annual_rent", "kaltmiete_jahr", "miete/jahr"],
        "lease_start": ["mietbeginn", "lease_start", "vertragsbeginn"],
        "lease_end": ["mietende", "lease_end", "vertragsende"],
        "tenant_type": ["mietertyp", "tenant_type", "kategorie"],
        "creditworthiness": ["bonität", "bonitaet", "creditworthiness", "rating"],
    }

    col_indices = {}
    for field, synonyms in col_map.items():
        for idx, header in enumerate(header_row):
            if any(syn in header for syn in synonyms):
                col_indices[field] = idx
                break

    tenants = []
    for row in rows[header_idx + 1:]:
        if not any(cell is not None for cell in row):
            continue
        tenant = {}
        for field, idx in col_indices.items():
            if idx < len(row):
                tenant[field] = _safe_value(row[idx])
        if tenant:
            tenants.append(tenant)

    return tenants


def parse_contract_overview(file_bytes: bytes) -> list:
    """
    Parse the 'Vertragsübersicht' sheet.
    Returns list of contract dicts.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    sheet_name = None
    for name in wb.sheetnames:
        if "vertrag" in name.lower() or "contract" in name.lower():
            sheet_name = name
            break
    if sheet_name is None and len(wb.sheetnames) > 2:
        sheet_name = wb.sheetnames[2]
    elif sheet_name is None and wb.sheetnames:
        sheet_name = wb.sheetnames[0]

    if sheet_name is None:
        return []

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = None
    header_idx = 0
    for i, row in enumerate(rows):
        if any(cell is not None for cell in row):
            header_row = [str(c).strip().lower() if c else "" for c in row]
            header_idx = i
            break

    if header_row is None:
        return []

    contracts = []
    for row in rows[header_idx + 1:]:
        if not any(cell is not None for cell in row):
            continue
        contract = {}
        for idx, header in enumerate(header_row):
            if idx < len(row) and row[idx] is not None:
                contract[header] = _safe_value(row[idx])
        if contract:
            contracts.append(contract)

    return contracts


def parse_floor_plan(file_bytes: bytes) -> list:
    """
    Parse the 'Flächenübersicht' sheet.
    Returns list of area dicts.
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

    sheet_name = None
    for name in wb.sheetnames:
        if "fläche" in name.lower() or "flaeche" in name.lower() or "floor" in name.lower():
            sheet_name = name
            break
    if sheet_name is None and len(wb.sheetnames) > 3:
        sheet_name = wb.sheetnames[3]
    elif sheet_name is None and wb.sheetnames:
        sheet_name = wb.sheetnames[0]

    if sheet_name is None:
        return []

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    header_row = None
    header_idx = 0
    for i, row in enumerate(rows):
        if any(cell is not None for cell in row):
            header_row = [str(c).strip().lower() if c else "" for c in row]
            header_idx = i
            break

    if header_row is None:
        return []

    areas = []
    for row in rows[header_idx + 1:]:
        if not any(cell is not None for cell in row):
            continue
        area = {}
        for idx, header in enumerate(header_row):
            if idx < len(row) and row[idx] is not None:
                area[header] = _safe_value(row[idx])
        if area:
            areas.append(area)

    return areas


def create_excel_template() -> bytes:
    """
    Create a multi-sheet Excel template with headers and example data.
    Returns bytes of the Excel file.
    """
    wb = openpyxl.Workbook()

    # Styles
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1B3A6B", end_color="1B3A6B", fill_type="solid")
    subheader_font = Font(bold=True, color="1B3A6B", size=10)
    example_fill = PatternFill(start_color="EBF0F8", end_color="EBF0F8", fill_type="solid")
    border_side = Side(style="thin", color="B0BEC5")
    cell_border = Border(
        left=border_side, right=border_side,
        top=border_side, bottom=border_side
    )
    center_align = Alignment(horizontal="center", vertical="center")

    def style_header_row(ws, row_num, num_cols):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = cell_border

    def style_data_row(ws, row_num, num_cols, is_example=True):
        fill = example_fill if is_example else PatternFill()
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_num, column=col)
            if is_example:
                cell.fill = fill
            cell.border = cell_border
            cell.alignment = Alignment(vertical="center")

    # ── Sheet 1: Objektstammdaten ──────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Objektstammdaten"
    ws1.column_dimensions["A"].width = 30
    ws1.column_dimensions["B"].width = 40

    ws1["A1"] = "Feld"
    ws1["B1"] = "Wert"
    style_header_row(ws1, 1, 2)

    property_data = [
        ("Objektname", "Bürogebäude Musterstraße"),
        ("Adresse", "Musterstraße 1"),
        ("Stadt", "Frankfurt am Main"),
        ("PLZ", "60311"),
        ("Objekttyp", "OFFICE"),
        ("Baujahr", 2005),
        ("Gesamtfläche (m²)", 2500.0),
        ("Grundstücksfläche (m²)", 1200.0),
        ("Etagen", 6),
        ("Einheiten", 12),
        ("Kaufpreis (€)", 8500000),
        ("Kaufdatum", "2024-01-15"),
    ]

    for i, (field, value) in enumerate(property_data, start=2):
        ws1.cell(row=i, column=1, value=field).font = subheader_font
        ws1.cell(row=i, column=1).border = cell_border
        ws1.cell(row=i, column=2, value=value)
        style_data_row(ws1, i, 2)

    ws1.row_dimensions[1].height = 25

    # ── Sheet 2: Mieterliste ───────────────────────────────────────────────
    ws2 = wb.create_sheet("Mieterliste")
    tenant_headers = [
        "Mietername", "Einheit", "Fläche (m²)", "Monatsmiete (€)",
        "Jahresmiete (€)", "Mietbeginn", "Mietende", "Mietertyp", "Bonität"
    ]
    col_widths = [25, 12, 14, 18, 16, 14, 14, 14, 10]
    for i, (h, w) in enumerate(zip(tenant_headers, col_widths), start=1):
        ws2.cell(row=1, column=i, value=h)
        ws2.column_dimensions[get_column_letter(i)].width = w

    style_header_row(ws2, 1, len(tenant_headers))

    tenant_examples = [
        ("Musterfirma GmbH", "EG-01", 500.0, 8500.0, 102000.0, "2022-01-01", "2026-12-31", "ANCHOR", "A"),
        ("Beispiel AG", "1.OG-01", 350.0, 5775.0, 69300.0, "2021-06-01", "2025-05-31", "STANDARD", "B"),
        ("Kleine UG", "2.OG-01", 120.0, 1800.0, 21600.0, "2023-03-01", "2024-02-29", "SMALL", "B"),
        ("Anker Retail GmbH", "EG-02", 800.0, 12000.0, 144000.0, "2020-01-01", "2030-12-31", "ANCHOR", "A"),
        ("Service KG", "3.OG-01", 200.0, 3200.0, 38400.0, "2022-09-01", "2025-08-31", "STANDARD", "C"),
    ]

    for row_idx, row_data in enumerate(tenant_examples, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws2.cell(row=row_idx, column=col_idx, value=value)
        style_data_row(ws2, row_idx, len(tenant_headers))

    ws2.row_dimensions[1].height = 25

    # ── Sheet 3: Vertragsübersicht ─────────────────────────────────────────
    ws3 = wb.create_sheet("Vertragsübersicht")
    contract_headers = [
        "Mietername", "Einheit", "Vertragsbeginn", "Vertragsende",
        "Monatliche Miete (€)", "Jährliche Miete (€)", "Indexierung (%)",
        "Optionen", "Kaution (€)", "Bemerkungen"
    ]
    contract_widths = [25, 12, 16, 16, 20, 18, 16, 20, 14, 30]
    for i, (h, w) in enumerate(zip(contract_headers, contract_widths), start=1):
        ws3.cell(row=1, column=i, value=h)
        ws3.column_dimensions[get_column_letter(i)].width = w

    style_header_row(ws3, 1, len(contract_headers))

    contract_examples = [
        ("Musterfirma GmbH", "EG-01", "2022-01-01", "2026-12-31", 8500.0, 102000.0, 2.0, "1x5 Jahre", 25500.0, ""),
        ("Beispiel AG", "1.OG-01", "2021-06-01", "2025-05-31", 5775.0, 69300.0, 1.5, "Keine", 17325.0, "Staffelmiete ab 2023"),
        ("Kleine UG", "2.OG-01", "2023-03-01", "2024-02-29", 1800.0, 21600.0, 0.0, "Keine", 5400.0, "Kurzvertrag"),
        ("Anker Retail GmbH", "EG-02", "2020-01-01", "2030-12-31", 12000.0, 144000.0, 2.5, "2x5 Jahre", 36000.0, "Langzeitmieter"),
        ("Service KG", "3.OG-01", "2022-09-01", "2025-08-31", 3200.0, 38400.0, 1.8, "1x3 Jahre", 9600.0, ""),
    ]

    for row_idx, row_data in enumerate(contract_examples, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws3.cell(row=row_idx, column=col_idx, value=value)
        style_data_row(ws3, row_idx, len(contract_headers))

    ws3.row_dimensions[1].height = 25

    # ── Sheet 4: Flächenübersicht ──────────────────────────────────────────
    ws4 = wb.create_sheet("Flächenübersicht")
    floor_headers = [
        "Etage", "Einheit", "Nutzungsart", "Fläche BGF (m²)",
        "Fläche NUF (m²)", "Fläche NGF (m²)", "Vermietet", "Mieter"
    ]
    floor_widths = [12, 14, 18, 16, 16, 16, 12, 25]
    for i, (h, w) in enumerate(zip(floor_headers, floor_widths), start=1):
        ws4.cell(row=1, column=i, value=h)
        ws4.column_dimensions[get_column_letter(i)].width = w

    style_header_row(ws4, 1, len(floor_headers))

    floor_examples = [
        ("EG", "EG-01", "Einzelhandel", 550.0, 500.0, 520.0, "Ja", "Musterfirma GmbH"),
        ("EG", "EG-02", "Einzelhandel", 880.0, 800.0, 840.0, "Ja", "Anker Retail GmbH"),
        ("1.OG", "1.OG-01", "Büro", 390.0, 350.0, 370.0, "Ja", "Beispiel AG"),
        ("1.OG", "1.OG-02", "Büro", 390.0, 350.0, 370.0, "Nein", "Leerstand"),
        ("2.OG", "2.OG-01", "Büro", 134.0, 120.0, 128.0, "Ja", "Kleine UG"),
        ("3.OG", "3.OG-01", "Büro", 223.0, 200.0, 212.0, "Ja", "Service KG"),
    ]

    for row_idx, row_data in enumerate(floor_examples, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws4.cell(row=row_idx, column=col_idx, value=value)
        style_data_row(ws4, row_idx, len(floor_headers))

    ws4.row_dimensions[1].height = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
